#!/usr/bin/env python3
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parent.parent
docx_path = root / "Lab3_AttackLab_Report.docx"
pdf_path = root / "Lab3_AttackLab_Report.pdf"

doc = Document(docx_path)
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print("headings:", len(headings))
for h in headings:
    print(" -", h)
assert not any("`" in h for h in headings), "literal backtick left in a heading"
text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            text += "\n" + cell.text
for needle in [
    "2.5  call、prologue、leave、ret 的作用",
    "8.1  dark-log 的栈布局",
    "8.2  利用 len 把 memcpy 的目标“抬高”",
    "payload_canary.py",
    "You have successfully detonated the bomb",
    "Hope A will be your grade",
]:
    assert needle in text, f"missing: {needle}"
print("docx paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
raw = pdf_path.read_bytes()
print("pdf bytes:", len(raw), "header:", raw[:8])
assert raw.startswith(b"%PDF-")
print("QA OK")
